"""Independent unit tests for xlsx2md."""
from pathlib import Path

from openpyxl import Workbook

from xlsx2md import (
    format_multi_note_text,
    get_image_row_and_col,
    process_single_sheet,
    remove_sheet_dir_if_exists,
)


def test_sample_conversion(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sample"
    sheet["B1"] = "标题"
    sheet["B2"] = "第二节"
    output_dir = tmp_path / "Sample"

    assert process_single_sheet(sheet, str(output_dir), sheet.title) is True
    result = (output_dir / "Sample.md").read_text(encoding="utf-8")
    assert "### 1.标题" in result
    assert "### 2.第二节" in result


def test_reject_sheet_when_column_a_contains_text(tmp_path):
    sheet = Workbook().active
    sheet["A1"] = "不允许的文字"
    sheet["B1"] = "标题"
    output_dir = tmp_path / "Skip"
    assert process_single_sheet(sheet, str(output_dir), "Skip") is False
    assert not output_dir.exists()


def test_reject_sheet_when_first_valid_row_is_missing(tmp_path):
    sheet = Workbook().active
    output_dir = tmp_path / "Empty"
    assert process_single_sheet(sheet, str(output_dir), "Empty") is False
    assert not output_dir.exists()


def test_note_formatting():
    assert "注1：" in format_multi_note_text("第一条\n第二条")
    assert "注2：" in format_multi_note_text("第一条\n第二条")
    assert format_multi_note_text("  \n  ") == ""


def test_remove_existing_sheet_directory(tmp_path):
    output_dir = tmp_path / "old"
    output_dir.mkdir()
    (output_dir / "stale.md").write_text("old", encoding="utf-8")
    remove_sheet_dir_if_exists(str(output_dir))
    assert not output_dir.exists()


def test_image_anchor_without_position():
    class EmptyImage:
        pass

    assert get_image_row_and_col(EmptyImage()) == (None, None)
