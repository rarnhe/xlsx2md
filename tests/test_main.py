"""Tests for the repository-root xlsx2md package."""
import os
from pathlib import Path
import subprocess
import sys

from openpyxl import Workbook

# Checkout directory itself is the package; its parent must be importable.
REPO_ROOT = Path(__file__).resolve().parents[1]
PARENT_DIR = REPO_ROOT.parent
if str(PARENT_DIR) not in sys.path:
    sys.path.insert(0, str(PARENT_DIR))

from xlsx2md.__main__ import (  # noqa: E402
    format_multi_note_text,
    get_default_excel_path,
    get_image_row_and_col,
    process_single_sheet,
    remove_sheet_dir_if_exists,
)


def test_sample_conversion(tmp_path):
    sheet = Workbook().active
    sheet.title = "Sample"
    sheet["B1"] = "标题"
    sheet["B2"] = "第二节"
    output_dir = tmp_path / "Sample"
    assert process_single_sheet(sheet, str(output_dir), sheet.title)
    text = (output_dir / "Sample.md").read_text(encoding="utf-8")
    assert "### 1.标题" in text
    assert "### 2.第二节" in text


def test_reject_sheet_with_column_a_text(tmp_path):
    sheet = Workbook().active
    sheet["A1"] = "禁止"
    sheet["B1"] = "标题"
    output_dir = tmp_path / "Skip"
    assert process_single_sheet(sheet, str(output_dir), "Skip") is False
    assert not output_dir.exists()


def test_reject_empty_sheet(tmp_path):
    sheet = Workbook().active
    output_dir = tmp_path / "Empty"
    assert process_single_sheet(sheet, str(output_dir), "Empty") is False
    assert not output_dir.exists()


def test_note_formatting():
    result = format_multi_note_text("第一条\n第二条")
    assert "注1：" in result and "注2：" in result
    assert format_multi_note_text("  \n  ") == ""


def test_remove_existing_sheet_directory(tmp_path):
    output_dir = tmp_path / "old"
    output_dir.mkdir()
    (output_dir / "stale.md").write_text("old", encoding="utf-8")
    remove_sheet_dir_if_exists(str(output_dir))
    assert not output_dir.exists()


def test_image_anchor_without_position():
    class EmptyImage:
        pass

    assert get_image_row_and_col(EmptyImage()) == (None, None)


def test_module_entrypoint_help():
    env = os.environ.copy()
    env["PYTHONIOENCODING"] = "utf-8"
    env["PYTHONPATH"] = str(PARENT_DIR)
    result = subprocess.run(
        [sys.executable, "-m", "xlsx2md", "--help"],
        cwd=PARENT_DIR,
        env=env,
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=False,
    )
    assert result.returncode == 0, result.stderr
    assert "sheet" in result.stdout


def test_default_workbook_path():
    assert Path(get_default_excel_path()).resolve() == REPO_ROOT / "xlsx2md.xlsx"
